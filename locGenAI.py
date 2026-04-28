import pandas as pd
import warnings
import json
import time
import sys
import os
from openai import OpenAI
from openpyxl import load_workbook

warnings.filterwarnings("ignore", category=UserWarning, module='urllib3')

# ===== CLI =====
if len(sys.argv) < 2 or sys.argv[1] in ("-h", "--help"):
    print("""
用法:   python3 locGenAI.py <Excel文件> [--test]

示例:   python3 locGenAI.py 国际化多语言.xlsx
        python3 locGenAI.py 国际化多语言.xlsx --test
        python3 locGenAI.py 国际化多语言.xlsx --swift    (同时生成 Swift 代码)

环境变量:
  DEEPSEEK_API_KEY   (必填) API 密钥
  DEEPSEEK_BASE_URL  (可选) API 地址，默认 https://api.deepseek.com
  DEEPSEEK_MODEL     (可选) 模型名，默认 deepseek-v4-flash
""")
    sys.exit(0)

excel_file = sys.argv[1]
test_mode = "--test" in sys.argv
swift_mode = "--swift" in sys.argv

if not os.path.exists(excel_file):
    print(f"❌ 找不到文件: {excel_file}")
    sys.exit(1)

# ===== 配置 =====
API_KEY = os.environ.get("DEEPSEEK_API_KEY", "")
API_BASE = os.environ.get("DEEPSEEK_BASE_URL", "https://api.deepseek.com")
API_MODEL = os.environ.get("DEEPSEEK_MODEL", "deepseek-v4-flash")
CACHE_FILE = "translation_cache.json"

if not API_KEY:
    print("❌ 请设置环境变量 DEEPSEEK_API_KEY")
    print("   export DEEPSEEK_API_KEY=sk-xxx")
    sys.exit(1)

client = OpenAI(api_key=API_KEY, base_url=API_BASE)

# ===== 持久化缓存 =====
def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_cache(cache):
    try:
        with open(CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(cache, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

translation_cache = load_cache()

def translate_to_key(text):
    if not text or pd.isna(text):
        return ""

    if text in translation_cache:
        return translation_cache[text]

    prompt = f"""
请将下面中文文案概括成最多三个英文单词，用下划线连接，生成 iOS Key：
只保留核心意思，不要输出完整句子，不要多余解释：
\"{text}\"
"""
    try:
        response = client.chat.completions.create(
            model=API_MODEL,
            messages=[
                {"role": "system", "content": "You are a helpful assistant"},
                {"role": "user", "content": prompt}
            ],
            stream=False
        )
        key_text = response.choices[0].message.content.strip()
        key = "_".join(key_text.lower().replace("/", "_").split()[:3])
        time.sleep(0.6)
    except Exception:
        print(f"\n  ⚠️  AI失败，回退: {text[:40]}")
        key = "_".join(text.strip().lower().replace("/", "_").split()[:3])

    translation_cache[text] = key
    return key


# ===== 进度条 =====
try:
    from tqdm import tqdm
    HAS_TQDM = True
except ImportError:
    HAS_TQDM = False


# ===== Excel 处理 =====
xls = pd.ExcelFile(excel_file)
sheet_count = len(xls.sheet_names)
sheet_module_keys = {}  # sheet_name → module_key，供后续 Swift 生成用

for sheet_index, sheet_name in enumerate(xls.sheet_names, 1):
    print(f"\n[{sheet_index}/{sheet_count}] 处理 Sheet: '{sheet_name}'")

    df = pd.read_excel(excel_file, sheet_name=sheet_name, header=1)
    df.columns = [str(c).strip() for c in df.columns]

    # ===== 自动检测列 =====
    text_col = next((c for c in df.columns if "中文" in c and "繁体" not in c), None)

    data_cols = [c for c in df.columns if "截图" not in c and "编码" not in c and "备注" not in c]
    page_col = data_cols[0] if len(data_cols) >= 1 and data_cols[0] != text_col else None
    candidate_ui = data_cols[1] if len(data_cols) >= 2 else None
    ui_col = candidate_ui if candidate_ui and candidate_ui != text_col else None

    # 语言列
    lang_cols = {}
    _lang_rules = [
        (["英文"], "en"),
        (["日语", "ja-JP", "ja"], "ja"),
        (["印度尼西亚", "id-ID"], "id"),
        (["繁体中文"], "zh-Hant"),
    ]
    for keywords, locale in _lang_rules:
        for kw in keywords:
            col = next((c for c in df.columns if kw in c), None)
            if col:
                lang_cols[locale] = col
                break
    lang_cols["zh-Hans"] = text_col

    # 编码列 = key 列，没有则跳过此 Sheet
    code_col = next((c for c in df.columns if "编码" in c), None)
    if not code_col:
        print(f"  ⚠️  无「编码」列，跳过此 Sheet")
        continue

    if test_mode:
        print(f"  🔍 列检测: page_col='{page_col}' | ui_col='{ui_col}' | text_col='{text_col}' | code_col='{code_col}' | 语言列: {list(lang_cols.keys())}")

    if not text_col:
        print("⚠️ 找不到中文列，跳过")
        continue

    # 填充合并单元格
    if page_col:
        df[page_col] = df[page_col].ffill()
    if ui_col:
        df[ui_col] = df[ui_col].ffill()

    # 编码列即 key 列，不再新增
    key_from_excel = pd.notna(df[code_col]).any() if code_col in df.columns else False

    sheet_rows = []
    total_rows = len(df)
    existing_keys = set()
    module_key = translate_to_key(sheet_name)
    sheet_module_keys[sheet_name] = module_key

    test_rows = df.head(5) if test_mode else None
    rows_iter = test_rows.iterrows() if test_mode else df.iterrows()

    # 进度条
    if not test_mode:
        # 先统计需要处理的行数
        pending = sum(1 for _, row in rows_iter
                      if not (pd.notna(row.get("key", "")) and str(row.get("key", "")).strip()))
        rows_iter = df.iterrows()  # 重新迭代
        pbar = tqdm(total=pending, desc=f"  {sheet_name}", unit="行", ncols=80) if HAS_TQDM else None
        processed = 0
    else:
        pbar = None

    for idx, row in rows_iter:
        row_number = idx + 1

        # 跳过已有 key（编码列有值）
        key_val = row.get(code_col, "")
        has_existing_key = pd.notna(key_val) and str(key_val).strip() != ""
        if not test_mode and has_existing_key:
            continue

        # 跳过翻译不完整的行（任一语言列为空则整行跳过）
        incomplete = False
        for col in lang_cols.values():
            val = row.get(col, "")
            if not (pd.notna(val) and str(val).strip()):
                incomplete = True
                break
        if incomplete:
            if test_mode:
                print(f"  ⏭️  行{row_number}: 翻译不完整，正式运行将跳过")
            continue

        row_texts = {}
        for loc, col in lang_cols.items():
            val = row.get(col, "")
            text = str(val).strip() if pd.notna(val) else ""
            if text:
                row_texts[loc] = text
        text_zh_hans = row_texts.get("zh-Hans", "")
        page_name = row.get(page_col, "") if page_col else ""
        ui_name = row.get(ui_col, "") if ui_col else ""
        page_key = translate_to_key(page_name)
        ui_key = translate_to_key(ui_name)
        text_key = translate_to_key(text_zh_hans)

        key_parts = [page_key, ui_key, text_key]
        base_key = ".".join([p for p in key_parts if p])

        counter = 1
        key = base_key
        while key in existing_keys:
            counter += 1
            key = f"{base_key}_{counter}"
        existing_keys.add(key)

        if test_mode:
            print(f"  ┌─ 行{row_number}:")
            print(f"  │   page='{page_name}' → {page_key}")
            print(f"  │   ui='{ui_name}' → {ui_key}")
            print(f"  │   text='{str(text_zh_hans)[:50]}' → {text_key}")
            lang_preview = " | ".join(f"{loc}={str(row_texts.get(loc, ''))[:20]}"
                                      for loc in lang_cols if row_texts.get(loc))
            print(f"  │   🌐 {lang_preview}")
            print(f"  └─ KEY: {key}")
        else:
            df.at[idx, code_col] = key
            if pbar:
                pbar.update(1)
            processed = processed + 1 if pbar else 0

        sheet_rows.append({
            "key": key,
            "texts": row_texts
        })

    if pbar:
        pbar.close()

    if test_mode:
        text_count = df[text_col].dropna().apply(lambda x: str(x).strip() != "").sum()
        print(f"  📊 总行数: {total_rows} | 有中文文本: {text_count}")
        if page_col:
            print(f"     page 唯一值: {df[page_col].dropna().nunique()}")
        if ui_col:
            print(f"     ui 唯一值:   {df[ui_col].dropna().nunique()}")
        continue

    # ===== 写回 Excel（编码列） =====
    wb = load_workbook(excel_file)
    ws = wb[sheet_name]
    header_row = 2

    code_col_idx = list(df.columns).index(code_col) + 1  # 0-indexed → 1-indexed

    for i, (_, row_data) in enumerate(df.iterrows(), start=3):
        ws.cell(row=i, column=code_col_idx, value=row_data[code_col])

    wb.save(excel_file)

    # ===== 合并写入 xcstrings =====
    file_name = f"{sheet_name}.xcstrings"
    if os.path.exists(file_name):
        with open(file_name, "r", encoding="utf-8") as f:
            string_catalog = json.load(f)
        old_count = len(string_catalog.get("strings", {}))
    else:
        string_catalog = {"sourceLanguage": "en", "strings": {}, "version": "1.1"}
        old_count = 0

    for row in sheet_rows:
        localizations = {}
        for loc, text in row["texts"].items():
            localizations[loc] = {"stringUnit": {"state": "translated", "value": text}}
        string_catalog["strings"][row["key"]] = {
            "extractionState": "manual",
            "localizations": localizations
        }

    with open(file_name, "w", encoding="utf-8") as f:
        json.dump(string_catalog, f, ensure_ascii=False, indent=2)

    total = len(string_catalog["strings"])
    if old_count:
        print(f"  📄 已更新 {file_name}，原有 {old_count} 条 + 新增 {len(sheet_rows)} 条 = 共 {total} 条")
    else:
        print(f"  📄 已生成 {file_name}，共 {total} 条")

# 保存缓存
save_cache(translation_cache)
if translation_cache:
    print(f"\n💾 翻译缓存已保存 ({len(translation_cache)} 条) -> {CACHE_FILE}")

if test_mode:
    print("\n🧪 测试完成（未写 Excel / 未生成 xcstrings）")
else:
    # ===== 生成 Swift 代码 =====
    if swift_mode:
        from gen_swift_l10n import xcstrings_to_swift
        print("\n--- Swift 代码生成 ---")
        for sheet_name in xls.sheet_names:
            xc_file = f"{sheet_name}.xcstrings"
            if os.path.exists(xc_file):
                swift_name = sheet_module_keys.get(sheet_name, "")
                swift_name = swift_name.replace("_", " ").title().replace(" ", "")
                xcstrings_to_swift(xc_file, swift_name=swift_name)
        print("--- Swift 代码生成完成 ---")

    print("\n✅ 全部完成")
